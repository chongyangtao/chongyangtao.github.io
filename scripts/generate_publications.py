from __future__ import annotations

import argparse
import html
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ALLOWED_TAGS = {
    "reward-model",
    "context-engineering",
    "multi-party-conversation",
    "generation",
    "information-retrieval",
}
REQUIRED = {"id", "section", "year", "title", "authors", "paper_url", "tags", "enabled"}
SAFE_HTML_TAGS = {"strong", "em", "b", "u", "sup", "br", "span", "font", "img"}


def text(value) -> str:
    return "" if value is None else str(value).strip()


def number(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def validate_url(value: str, field: str, row: int, warnings: list[str]) -> None:
    parsed = urlparse(value)
    if value == "XXX" or "XXX" in value:
        warnings.append(f"row {row}: {field} contains placeholder URL {value!r}")
    elif parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError(f"row {row}: {field} must be an http(s) URL")


def validate_html(value: str, field: str, row: int) -> None:
    for tag in re.findall(r"</?([A-Za-z][\w-]*)", value):
        if tag.lower() not in SAFE_HTML_TAGS:
            raise ValueError(f"row {row}: {field} contains unsupported HTML tag <{tag}>")
    if re.search(r"\son\w+\s*=|javascript:", value, re.I):
        raise ValueError(f"row {row}: {field} contains unsafe HTML")


def compact_publication_divs(markup: str) -> str:
    pattern = re.compile(
        r"[ \t]*<div class=\"publication-(?:title|authors|venue|links)\">.*?</div>[ \t]*\n?",
        re.DOTALL,
    )

    def compact(match: re.Match[str]) -> str:
        fragment = BeautifulSoup(match.group(0), "html.parser")
        element = fragment.find("div")
        if not element:
            return match.group(0)
        one_line = re.sub(r"\s+", " ", str(element)).strip()
        return "        " + one_line + "\n"

    return pattern.sub(compact, markup)


def render_static_html(records: list[dict]) -> None:
    page_path = ROOT / "publications.html"
    soup = BeautifulSoup(page_path.read_text(encoding="utf-8"), "html.parser")
    container = soup.select_one(".publication-lists")
    if not container:
        raise ValueError("publications.html must contain a .publication-lists container")

    grouped = OrderedDict()
    for record in records:
        grouped.setdefault(record["section"], []).append(record)

    container.clear()
    for section, items in grouped.items():
        year_group = soup.new_tag("section", attrs={"class": "publication-year"})
        heading = soup.new_tag("strong")
        heading.string = section
        year_group.append(heading)
        listing = soup.new_tag("ul", attrs={"class": "publication-list"})
        for record in items:
            item = soup.new_tag("li", attrs={
                "class": "publication-item",
                "data-tags": " ".join(record["tags"]),
                "data-id": record["id"],
            })
            title = soup.new_tag("div", attrs={"class": "publication-title"})
            title_link = soup.new_tag("a", href=record["paper_url"], target="_blank", rel="noopener noreferrer")
            title_link.string = record["title"]
            title.append(title_link)
            item.append(title)
            authors = soup.new_tag("div", attrs={"class": "publication-authors"})
            authors.append(BeautifulSoup(record.get("authors_html") or html.escape(record["authors"]), "html.parser"))
            item.append(authors)
            venue = soup.new_tag("div", attrs={"class": "publication-venue"})
            venue.append(BeautifulSoup(record.get("venue_html") or html.escape(record["venue"]), "html.parser"))
            item.append(venue)
            if record.get("code_url"):
                links = soup.new_tag("div", attrs={"class": "publication-links"})
                link = soup.new_tag("a", href=record["code_url"], target="_blank", rel="noopener noreferrer")
                link.string = record.get("code_label") or "[Link]"
                links.append(link)
                item.append(links)
            listing.append(item)
        year_group.append(listing)
        container.append("\n")
        container.append(BeautifulSoup(year_group.prettify(formatter="html"), "html.parser"))
        container.append("\n")

    loading = soup.select_one(".publication-loading")
    if loading:
        loading.decompose()
    page_path.write_text(compact_publication_divs(soup.prettify(formatter="html")), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate publication JSON from Excel.")
    parser.add_argument("--input", type=Path, default=ROOT / "data/publications.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "files/publications.json")
    args = parser.parse_args()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    if "Publications" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a Publications sheet")
    sheet = workbook["Publications"]
    rows = sheet.iter_rows(values_only=True)
    headers = {text(value): index for index, value in enumerate(next(rows, ())) if text(value)}
    missing = REQUIRED - headers.keys()
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")

    records = []
    seen_ids = set()
    warnings = []
    for row_number, values in enumerate(rows, start=2):
        row = {key: values[index] if index < len(values) else None for key, index in headers.items()}
        if not any(value is not None and text(value) for value in values):
            continue
        record = {key: text(row.get(key)) for key in headers}
        record["year"] = number(row.get("year"), 0)
        record["month"] = number(row.get("month"), 0)
        record["sort_order"] = number(row.get("sort_order"), row_number)
        record["enabled"] = str(row.get("enabled", "")).strip().lower() not in {"false", "0", "no", "n"}
        if not record["id"] or record["id"] in seen_ids:
            raise ValueError(f"row {row_number}: id is empty or duplicated")
        seen_ids.add(record["id"])
        for field in ("section", "title", "authors", "paper_url", "tags"):
            if not record[field]:
                raise ValueError(f"row {row_number}: {field} is required")
        validate_url(record["paper_url"], "paper_url", row_number, warnings)
        if record.get("code_url"):
            try:
                validate_url(record["code_url"], "code_url", row_number, warnings)
            except ValueError as error:
                warnings.append(str(error))
        tags = record["tags"].split()
        unknown = set(tags) - ALLOWED_TAGS
        if unknown:
            raise ValueError(f"row {row_number}: unknown tags: {', '.join(sorted(unknown))}")
        record["tags"] = tags
        for field in ("authors_html", "venue_html"):
            validate_html(record.get(field, ""), field, row_number)
        records.append(record)

    records.sort(key=lambda item: (0 if item["section"].lower() == "preprint" else 1, -item["year"], item["month"], item["sort_order"]))
    output = {
        "version": 1,
        "filters": [
            {"id": "all", "label": "All"},
            {"id": "reward-model", "label": "Reward Model"},
            {"id": "context-engineering", "label": "Context Engineering"},
            {"id": "multi-party-conversation", "label": "Multi-party Conversation"},
            {"id": "generation", "label": "Generation"},
            {"id": "information-retrieval", "label": "Information Retrieval"},
        ],
        "publications": [record for record in records if record["enabled"]],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    render_static_html(output["publications"])
    for warning in warnings:
        print(f"warning: {warning}", file=sys.stderr)
    print(f"Wrote {len(output['publications'])} publications to {args.output}")


if __name__ == "__main__":
    main()
